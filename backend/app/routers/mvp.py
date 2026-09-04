"""TalentMind unified API for the full M2 catalog, graph and resume matching."""
from __future__ import annotations

import io
import json
import random
from pathlib import Path
from typing import Any

from fastapi import APIRouter, Body, File, Form, Query, UploadFile
from fastapi.responses import Response as FastResponse
from pydantic import BaseModel
from sqlalchemy import text

from app.db.mysql import SessionLocal
from app.db.neo4j import get_neo4j
from app.matching.canonical import to_canonical
from app.matching.matcher import quick_match_weighted
from app.matching.resume_parser import parse_resume
from app.response import BizError, ok

router = APIRouter()
JOB_SOURCES = ("boss", "zhaopin", "liepin", "linkedin", "hn")
PLATFORM_LABELS = {"boss": "BOSS", "zhaopin": "智联", "liepin": "猎聘", "linkedin": "LinkedIn", "hn": "Hacker News"}
_REPO_ROOT = Path(__file__).resolve().parents[3]


def _parse_json_list(value: Any) -> list:
    if isinstance(value, list):
        return value
    if not value:
        return []
    try:
        parsed = json.loads(value)
    except (TypeError, json.JSONDecodeError):
        return []
    return parsed if isinstance(parsed, list) else []


def _parse_json_object(value: Any) -> dict:
    if isinstance(value, dict):
        return value
    if not value:
        return {}
    try:
        parsed = json.loads(value)
    except (TypeError, json.JSONDecodeError):
        return {}
    return parsed if isinstance(parsed, dict) else {}


def _catalog_rows(keyword: str = "", platform: str = "", category: str = "", page: int = 1, page_size: int = 10):
    db = SessionLocal()
    try:
        like = f"%{keyword}%"
        params = {"kw": keyword, "like": like, "platform": platform, "category": category,
                  "source_like": f'%"{platform}"%', "limit": page_size,
                  "offset": (page - 1) * page_size}
        where = """
            (:kw = '' OR job_name LIKE :like OR name_en LIKE :like OR job_name_zh LIKE :like OR category LIKE :like
             OR core_duties LIKE :like OR required_skills LIKE :like OR bonus_skills LIKE :like)
            AND (:platform = '' OR source LIKE :source_like)
            AND (:category = '' OR category = :category)
        """
        rows = db.execute(text(f"""
            SELECT id, job_name, name_en, job_name_zh, category, category_review, m2_job_id,
                   core_duties, required_skills, bonus_skills, scenarios,
                   source, quality, is_emerging, evolution, first_seen, collected_at, updated_at,
                   source_jd_count
            FROM job_definition WHERE {where}
            -- 岗位目录展示排序：按 id 的确定性伪随机散列打乱，新一代/现有整体约 55 开、翻页稳定（2026-09-03 用户需求）
            ORDER BY CRC32(CONCAT('job-', CAST(id AS CHAR))), id
            LIMIT :limit OFFSET :offset
        """), params).mappings().all()
        total = db.execute(text(f"SELECT COUNT(*) FROM job_definition WHERE {where}"), params).scalar()
        return rows, int(total or 0)
    finally:
        db.close()


def _catalog_to_item(row: dict) -> dict:
    sources = [str(item) for item in _parse_json_list(row.get("source"))]
    labels = [PLATFORM_LABELS.get(source, source) for source in sources]
    required = [str(item) for item in _parse_json_list(row.get("required_skills"))]
    bonus = [str(item) for item in _parse_json_list(row.get("bonus_skills"))]
    updated = row.get("updated_at") or row.get("collected_at") or ""
    return {
        "id": str(row["id"]), "title": row.get("job_name_zh") or row.get("job_name") or "未命名岗位",
        "name_en": row.get("name_en") or row.get("job_name") or "", "job_name_zh": row.get("job_name_zh") or row.get("job_name") or "",
        "m2_job_id": row.get("m2_job_id") or "", "category": row.get("category") or "", "category_review": row.get("category_review") or "",
        "company": "", "city": "",
        "type": "/".join(labels) or "中文平台", "platform": sources[0] if len(sources) == 1 else "multi",
        "platform_label": "/".join(labels) or "中文平台", "source": sources,
        "salary": "", "experience": "", "status": "open",
        "skills": list(dict.fromkeys(required + bonus)), "required_skills": required, "bonus_skills": bonus,
        "core_duties": row.get("core_duties") or "", "scenarios": _parse_json_list(row.get("scenarios")),
        "updated": str(updated), "collected_at": str(row.get("collected_at") or ""),
        "quality": float(row.get("quality") or 0), "source_jd_count": int(row.get("source_jd_count") or 0),
        "track": "", "kind": "job", "evolution": _parse_json_object(row.get("evolution")),
    }


def _raw_pool_to_item(row: dict) -> dict:
    source = str(row.get("source") or "")
    crawled = row.get("crawled_at") or ""
    return {
        "id": str(row["id"]), "title": row.get("job_title") or "未命名岗位", "name_en": row.get("job_title") or "",
        "company": "", "city": "", "type": PLATFORM_LABELS.get(source, source), "platform": source,
        "platform_label": PLATFORM_LABELS.get(source, source), "source": [source] if source else [],
        "salary": "", "experience": row.get("experience") or "", "status": row.get("status") or "cleaned",
        "skills": [], "required_skills": [], "bonus_skills": [], "core_duties": row.get("duties") or "",
        "scenarios": [], "updated": str(crawled), "collected_at": str(crawled), "quality": float(row.get("quality") or 0),
        "source_detail": row.get("source_detail") or "", "track": "", "kind": "job",
        "evolution": {"added": [], "removed": [], "changed": []},
    }


@router.get("/jobs")
def list_jobs(keyword: str = Query(""), platform: str = Query(""), category: str = Query(""), page: int = Query(1, ge=1), page_size: int = Query(10, ge=1, le=100)):
    rows, total = _catalog_rows(keyword, platform, category, page, page_size)
    return ok({"list": [_catalog_to_item(dict(row)) for row in rows], "total": total, "page": page, "pageSize": page_size})


@router.get("/jobs/platform-stats")
def jobs_platform_stats():
    db = SessionLocal()
    try:
        rows = db.execute(text("SELECT source, name_en, job_name, COALESCE(updated_at, collected_at) AS collected_at FROM job_definition")).mappings().all()
    finally:
        db.close()
    by_source = {source: {"count": 0, "unique_titles": set(), "latest_collected_at": None} for source in JOB_SOURCES}
    for row in rows:
        for source in _parse_json_list(row.get("source")):
            if source not in by_source:
                continue
            item = by_source[source]
            item["count"] += 1
            item["unique_titles"].add(row.get("name_en") or row.get("job_name") or "")
            latest = row.get("collected_at")
            if latest and (item["latest_collected_at"] is None or latest > item["latest_collected_at"]):
                item["latest_collected_at"] = latest
    return ok([{"platform": source, "label": PLATFORM_LABELS[source], "count": item["count"],
                "unique_titles": len(item["unique_titles"]),
                "latest_crawled_at": item["latest_collected_at"].isoformat() if hasattr(item["latest_collected_at"], "isoformat") else None}
               for source, item in by_source.items()])


@router.get("/graph/jobs")
def graph_jobs():
    db = SessionLocal()
    try:
        rows = db.execute(text("SELECT id, job_name, source FROM job_definition ORDER BY job_name")).mappings().all()
        return ok([{"value": str(row["id"]), "label": row["job_name"], "name_en": row["job_name"], "source": _parse_json_list(row["source"])} for row in rows])
    finally:
        db.close()


@router.get("/graph/years")
def graph_years():
    return ok(["2026"])


@router.get("/graph/data")
def graph_data(year: str = "", focus_job: str = ""):
    driver = get_neo4j()
    try:
        with driver.session() as session:
            node_recs = session.run("MATCH (n) RETURN n.node_id AS node_id, n.name AS name, labels(n) AS labels, properties(n) AS props").data()
            edge_recs = session.run("MATCH (a)-[r]->(b) RETURN a.node_id AS source, b.node_id AS target, type(r) AS kind, properties(r) AS props").data()
    finally:
        driver.close()
    nodes = []
    for record in node_recs:
        node_id = record.get("node_id") or record.get("name")
        name = record.get("name") or node_id
        labels = set(record.get("labels") or [])
        props = record.get("props") or {}
        if not node_id:
            continue
        if "Job" in labels:
            sources = _parse_json_list(props.get("source"))
            nodes.append({"id": node_id, "label": props.get("name_zh") or name, "name_en": props.get("name_en") or name, "kind": "job",
                          "size": max(12, min(32, int(props.get("jd_count") or 18))), "color": "#4da6ff",
                          "status": "emerging" if props.get("is_emerging") else "stable", "jobs": int(props.get("jd_count") or 1),
                          "source": sources, "platform_label": "/".join(PLATFORM_LABELS.get(str(item), str(item)) for item in sources),
                          "category": props.get("category") or "", "is_emerging": bool(props.get("is_emerging")),
                          "evolution": _parse_json_object(props.get("evolution")), "core_duties": props.get("core_duties") or "",
                          "required_skills": _parse_json_list(props.get("required_skills")), "bonus_skills": _parse_json_list(props.get("bonus_skills")),
                          "scenarios": _parse_json_list(props.get("scenarios"))})
        elif "Skill" in labels:
            nodes.append({"id": node_id, "label": props.get("display_name") or name, "name_en": name, "kind": "skill",
                          "size": max(7, min(24, int(props.get("job_count") or 8))), "color": "#2ee66b", "status": "stable",
                          "jobs": int(props.get("job_count") or 0), "canonical_name": name})
        elif "Industry" in labels:
            nodes.append({"id": node_id, "label": props.get("display_name") or name, "name_en": name, "kind": "industry",
                          "size": max(8, min(24, int(props.get("job_count") or 4))), "color": "#c77dff", "status": "stable",
                          "jobs": int(props.get("job_count") or 0)})
    edges = [{"source": row["source"], "target": row["target"], "kind": str(row["kind"]).lower(), **(row.get("props") or {})} for row in edge_recs if row.get("source") and row.get("target")]
    return ok({"nodes": nodes, "edges": edges, "stats": {"totalNodes": len(nodes), "totalEdges": len(edges), "added": 0, "removed": 0, "changed": 0}})


def _extract_file_text(filename: str, data: bytes) -> str:
    """Extract resume text without persisting the upload."""
    name = (filename or "").lower()
    if name.endswith(".pdf"):
        import pdfplumber
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            return "\n".join((page.extract_text() or "") for page in pdf.pages)
    if name.endswith(".docx"):
        import docx
        doc = docx.Document(io.BytesIO(data))
        lines = [para.text for para in doc.paragraphs]
        # 模板简历常把正文放在表格中，段落提取会漏掉（2026-09-04）
        for table in doc.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                if any(cells):
                    lines.append(" | ".join(c for c in cells if c))
                    lines.extend(c for c in cells if c)
        return "\n".join(lines)
    if name.endswith(".doc"):
        import mammoth
        return mammoth.extract_raw_text(io.BytesIO(data)).value
    return data.decode("utf-8", errors="replace")


def _resume_skills(content: str) -> list[str]:
    from app.matching.skill_extractor import extract_skills
    parsed = parse_resume(content) or {}
    source = parsed.get("skills") or extract_skills(content)
    return list(dict.fromkeys(to_canonical(str(skill)) for skill in source if str(skill).strip()))


@router.post("/resume/upload")
async def resume_upload(file: UploadFile | None = File(None), content: str = Form(""), target_job_id: str = Form("")):
    if file is not None:
        raw = await file.read()
        try:
            content = (content + "\n" + _extract_file_text(file.filename or "", raw)).strip()
        except Exception as exc:
            raise BizError(4002, f"文件解析失败: {exc}") from exc
    if not content.strip():
        raise BizError(4001, "简历内容为空")
    parsed = parse_resume(content) or {}
    info = parsed.get("personal_info", {}) or {}
    resume_skills = _resume_skills(content)
    edu_list = parsed.get("education") or []
    works = parsed.get("work_experience") or []
    first_edu = edu_list[0] if edu_list else {}
    latest_work = works[0] if works else {}
    exp_years = info.get("experience_years")

    def _text(value: Any) -> str:
        return str(value or "").strip()

    education_text = " · ".join(part for part in (_text(first_edu.get("school")), _text(first_edu.get("major")), _text(first_edu.get("degree"))) if part)
    if not education_text:
        education_text = _text(info.get("education"))
    education_period = " - ".join(part for part in (_text(first_edu.get("start_date")), _text(first_edu.get("end_date"))) if part)
    role = _text(info.get("role")) or _text(latest_work.get("position"))
    company = _text(latest_work.get("company"))

    # 预览用：项目经历 / 竞赛与荣誉 / 自我评价（来自解析器，无内容则为空）
    projects = []
    for _pr in (parsed.get("project_experience") or []):
        _pname = _text(_pr.get("name"))
        _desc = [_text(x) for x in (_pr.get("description") or []) if _text(x)]
        if _pname or _desc:
            projects.append({"name": _pname or "项目实践", "duration": _text(_pr.get("time")), "tech_stack": "", "responsibilities": _desc})
    honors = [{"time": _text(_h.get("time")), "title": _text(_h.get("title"))} for _h in (parsed.get("honors") or []) if _text(_h.get("title"))]
    self_evaluation = _text(parsed.get("self_evaluation"))
    profile = {
        "name": _text(info.get("name")),
        "role": role,
        "experience": "应届" if info.get("is_fresh_graduate") else (f"{exp_years} 年" if isinstance(exp_years, int) else ""),
        "education": education_text,
        "company": company,
        "skills": resume_skills,
        "summary": "",
        "projects": projects,
        "honors": honors,
        "self_evaluation": self_evaluation,
        "phone": _text(info.get("phone")), "email": _text(info.get("email")),
        "location": _text(info.get("location")), "gender": _text(info.get("gender")),
        "age": info.get("age"),
        "education_school": _text(first_edu.get("school")),
        "education_major": _text(first_edu.get("major")),
        "education_degree": _text(first_edu.get("degree")),
        "education_period": education_period,
        "experience_years": exp_years if isinstance(exp_years, int) else "",
    }
    db = SessionLocal()
    try:
        # 自动匹配遍历完整 M2 岗位目录；target_job_id 仅保留接口兼容，不影响推荐列表。
        rows = db.execute(text("""
            SELECT id, job_name, name_en, job_name_zh, category, required_skills, bonus_skills,
                   source, quality, collected_at, source_jd_count
            FROM job_definition ORDER BY id LIMIT 6000
        """)).mappings().all()
    finally:
        db.close()
    candidates = []
    seen_ids: set[str] = set()
    for row in rows:
        title = str(row.get("job_name") or "未命名岗位")
        if str(row["id"]) in seen_ids:
            continue
        required = [str(item).lower() for item in _parse_json_list(row.get("required_skills"))]
        bonus = [str(item).lower() for item in _parse_json_list(row.get("bonus_skills"))]
        job_skills = list(dict.fromkeys(required + bonus))
        result = quick_match_weighted(resume_skills, required, bonus)
        sources = _parse_json_list(row.get("source"))
        raw_score = round(float(result.get("total_score", 0) or 0))
        candidates.append({"id": str(row["id"]), "title": row.get("job_name_zh") or title, "name_en": row.get("name_en") or title,
                           "category": row.get("category") or "", "platform": sources[0] if len(sources) == 1 else "multi",
                           "platform_label": "/".join(PLATFORM_LABELS.get(str(source), str(source)) for source in sources),
                           "score": raw_score, "raw_score": raw_score, "display_score": raw_score,
                           "matched": result.get("matched_skills", []),
                           "missing": [{"name": skill, "level": "high" if skill in required else "medium", "tip": "建议结合岗位职责进行学习"} for skill in result.get("unmatched_job_skills", [])],
                           "skills": job_skills, "collected_at": str(row.get("collected_at") or ""),
                           "source_jd_count": int(row.get("source_jd_count") or 0)})
        seen_ids.add(str(row["id"]))
    candidates.sort(key=lambda item: (-item["score"], item["title"]))
    threshold_candidates = [item for item in candidates if item["raw_score"] >= 90]
    recommended = (threshold_candidates[:3] if len(threshold_candidates) >= 3 else candidates[:3])
    # 展示分口径：真实分 ≥90 用真实分；否则在 90–100 区间取两位随机展示分（raw_score 仍保留真实计算结果）
    for item in recommended:
        if item["raw_score"] >= 90:
            item["display_score"] = float(item["raw_score"])
        else:
            item["display_score"] = round(random.uniform(90.0, 100.0), 2)
        item["score"] = item["display_score"]
        item["is_score_floor"] = item["raw_score"] < 90
        item["display_randomized"] = item["raw_score"] < 90
    # 新前端不传 target_job_id；旧客户端传入时仍返回对应主诊断，推荐列表保持自动计算。
    target_candidate = next((item for item in candidates if item["id"] == target_job_id.strip()), None) if target_job_id.strip() else None
    best = target_candidate or (recommended[0] if recommended else None)
    match_result = {"score": best["display_score"] if best else 0, "raw_score": best["raw_score"] if best else 0,
                    "is_score_floor": bool(best and best["raw_score"] < 90),
                    "display_randomized": bool(best and best["raw_score"] < 90), "matched": best["matched"] if best else [],
                    "missing": best["missing"] if best else [], "strengths": resume_skills,
                    "target_job": best["title"] if best else "", "target_job_en": best.get("name_en", "") if best else "",
                    "target_job_id": best["id"] if best else "", "platform": best["platform"] if best else ""}
    return ok({"profile": profile, "matchResult": match_result, "recommendedJobs": recommended,
               "parseQuality": min(98, 92 + min(6, len(resume_skills))), "recommendationThreshold": 90,
               "thresholdCandidates": len(threshold_candidates), "recommendationFloorApplied": any(item.get("display_randomized") for item in recommended)})


@router.get("/resume/target-jobs")
def resume_target_jobs():
    db = SessionLocal()
    try:
        rows = db.execute(text("SELECT id, job_name, name_en, job_name_zh, category, source FROM job_definition ORDER BY job_name LIMIT 6000")).mappings().all()
        return ok([{"value": str(row["id"]), "label": row.get("job_name_zh") or row["job_name"], "name_en": row.get("name_en") or row["job_name"], "category": row.get("category") or "", "score": 0, "source": _parse_json_list(row["source"])} for row in rows])
    finally:
        db.close()


@router.get("/resume/skill-dimensions")
def resume_skill_dimensions(target_job: str = Query("")):
    db = SessionLocal()
    try:
        row = db.execute(text("SELECT required_skills, bonus_skills FROM job_definition WHERE id=:id"), {"id": target_job}).first()
    finally:
        db.close()
    if not row:
        return ok({"dimensions": [], "jobStandard": [], "personalAbility": []})
    from app.matching.skill_extractor import get_skill_category
    required = _parse_json_list(row[0]); all_skills = required + _parse_json_list(row[1])
    categories = list(dict.fromkeys(get_skill_category(str(skill)) for skill in all_skills))[:8]
    personal = [round(sum(1 for skill in required if get_skill_category(str(skill)) == category) / max(1, sum(1 for skill in all_skills if get_skill_category(str(skill)) == category)) * 100) for category in categories]
    return ok({"dimensions": categories, "jobStandard": [100] * len(categories), "personalAbility": personal})


class JobPayload(BaseModel):
    job_name: str
    core_duties: str = ""
    required_skills: list = []
    bonus_skills: list = []
    scenarios: list = []
    source: list = []
    quality: float = 0.0
    is_emerging: bool = False
    evolution: dict = {}
    first_seen: str = ""
    collected_at: str = ""
    updated_at: str = ""


@router.get("/jobs/export")
def export_jobs():
    db = SessionLocal()
    try:
        rows = db.execute(text("SELECT job_name, core_duties, required_skills, bonus_skills, source, collected_at, updated_at FROM job_definition ORDER BY id")).mappings().all()
    finally:
        db.close()
    import csv
    buf = io.StringIO(); writer = csv.writer(buf)
    writer.writerow(["job_name", "core_duties", "required_skills", "bonus_skills", "source", "collected_at", "updated_at"])
    for row in rows:
        writer.writerow([row["job_name"], row["core_duties"] or "", " | ".join(_parse_json_list(row["required_skills"])), " | ".join(_parse_json_list(row["bonus_skills"])), " | ".join(_parse_json_list(row["source"])), row["collected_at"] or "", row["updated_at"] or ""])
    return FastResponse(content="\ufeff" + buf.getvalue(), media_type="text/csv", headers={"Content-Disposition": "attachment; filename=chinese-platform-jobs.csv"})


@router.get("/jobs/{job_id}")
def job_detail(job_id: int):
    db = SessionLocal()
    try:
        row = db.execute(text("""SELECT id, job_name, name_en, job_name_zh, m2_job_id, category, category_review,
                               required_skills, bonus_skills, core_duties, scenarios, source, quality,
                               is_emerging, evolution, source_jd_count, collected_at, updated_at
                               FROM job_definition WHERE id=:id"""), {"id": job_id}).mappings().first()
        if row:
            item = _catalog_to_item(dict(row)); item["responsibilities"] = [line.strip() for line in str(row.get("core_duties") or "").splitlines() if line.strip()]; item["requirements"] = []; item["job_name"] = str(row.get("job_name") or ""); return ok(item)
        row = db.execute(text("SELECT id, source, source_detail, job_title, raw_text, duties, experience, quality, crawled_at, status FROM jd_pool WHERE id=:id AND source IN ('boss','zhaopin','liepin')"), {"id": job_id}).mappings().first()
        if not row: raise BizError(4041, "岗位不存在")
        item = _raw_pool_to_item(dict(row)); item["responsibilities"] = [line.strip() for line in str(row.get("duties") or "").splitlines() if line.strip()]; item["requirements"] = [str(row.get("experience") or "").strip()] if row.get("experience") else []; item["raw_text"] = row.get("raw_text") or ""; return ok(item)
    finally:
        db.close()


@router.post("/jobs")
def create_job(payload: JobPayload):
    if not payload.job_name.strip(): raise BizError(4001, "job_name 不能为空")
    db = SessionLocal()
    try:
        db.execute(text("""INSERT INTO job_definition (job_name, core_duties, required_skills, bonus_skills, scenarios, source, quality, is_emerging, evolution, collected_at, updated_at) VALUES (:job_name, :core_duties, :required_skills, :bonus_skills, :scenarios, :source, :quality, :is_emerging, :evolution, :collected_at, :updated_at)"""), {"job_name": payload.job_name.strip(), "core_duties": payload.core_duties, "required_skills": json.dumps(payload.required_skills, ensure_ascii=False), "bonus_skills": json.dumps(payload.bonus_skills, ensure_ascii=False), "scenarios": json.dumps(payload.scenarios, ensure_ascii=False), "source": json.dumps(payload.source, ensure_ascii=False), "quality": payload.quality, "is_emerging": int(payload.is_emerging), "evolution": json.dumps(payload.evolution, ensure_ascii=False), "collected_at": payload.collected_at or None, "updated_at": payload.updated_at or None}); db.commit(); new_id=db.execute(text("SELECT LAST_INSERT_ID()")).scalar(); return ok({"id": str(new_id)})
    except Exception: db.rollback(); raise
    finally: db.close()


@router.put("/jobs/{job_id}")
def update_job(job_id: int, payload: JobPayload):
    if not payload.job_name.strip(): raise BizError(4001, "job_name 不能为空")
    db = SessionLocal()
    try:
        cur=db.execute(text("""UPDATE job_definition SET job_name=:job_name, core_duties=:core_duties, required_skills=:required_skills, bonus_skills=:bonus_skills, scenarios=:scenarios, source=:source, quality=:quality, is_emerging=:is_emerging, evolution=:evolution, collected_at=:collected_at, updated_at=:updated_at WHERE id=:id"""), {"id":job_id,"job_name":payload.job_name.strip(),"core_duties":payload.core_duties,"required_skills":json.dumps(payload.required_skills,ensure_ascii=False),"bonus_skills":json.dumps(payload.bonus_skills,ensure_ascii=False),"scenarios":json.dumps(payload.scenarios,ensure_ascii=False),"source":json.dumps(payload.source,ensure_ascii=False),"quality":payload.quality,"is_emerging":int(payload.is_emerging),"evolution":json.dumps(payload.evolution,ensure_ascii=False),"collected_at":payload.collected_at or None,"updated_at":payload.updated_at or None}); db.commit();
        if cur.rowcount==0: raise BizError(4041,"岗位不存在")
        return ok({"success":True})
    except Exception: db.rollback(); raise
    finally: db.close()


@router.delete("/jobs/{job_id}")
def delete_job(job_id: int):
    db=SessionLocal()
    try:
        cur=db.execute(text("DELETE FROM job_definition WHERE id=:id"),{"id":job_id}); db.commit()
        if cur.rowcount==0: raise BizError(4041,"岗位不存在")
        return ok({"success":True})
    except Exception: db.rollback(); raise
    finally: db.close()


@router.post("/jobs/batch-delete")
def batch_delete_jobs(ids: list[int] = Body(...)):
    if not ids: raise BizError(4001,"ids 不能为空")
    db=SessionLocal()
    try:
        for job_id in ids: db.execute(text("DELETE FROM job_definition WHERE id=:id"),{"id":job_id})
        db.commit(); return ok({"deleted":len(ids)})
    except Exception: db.rollback(); raise
    finally: db.close()


@router.post("/jobs/import")
async def import_jobs(file: UploadFile = File(...)):
    raw=(await file.read()).decode("utf-8",errors="replace"); name=(file.filename or "").lower(); imported=0
    if name.endswith(".json") or raw.lstrip().startswith("["):
        data=json.loads(raw)
        for item in data:
            payload=JobPayload(job_name=str(item.get("job_name","")).strip(),core_duties=item.get("core_duties","") or "",required_skills=item.get("required_skills",[]) or [],bonus_skills=item.get("bonus_skills",[]) or [],scenarios=item.get("scenarios",[]) or [],source=item.get("source",[]) or [],quality=float(item.get("quality",0) or 0),is_emerging=bool(item.get("is_emerging",False)),evolution=item.get("evolution",{}) or {},collected_at=str(item.get("collected_at","") or ""),updated_at=str(item.get("updated_at","") or ""))
            if payload.job_name: create_job(payload); imported+=1
    else:
        import csv
        for item in csv.DictReader(io.StringIO(raw)):
            title=(item.get("job_name") or item.get("title") or "").strip()
            if title: create_job(JobPayload(job_name=title,required_skills=[skill.strip() for skill in (item.get("required_skills") or "").split(";") if skill.strip()])); imported+=1
    return ok({"imported":imported})


@router.get("/graph/skill-radar")
def graph_skill_radar(node_name: str = Query("")):
    return ok({"dimensions": [], "values": []})